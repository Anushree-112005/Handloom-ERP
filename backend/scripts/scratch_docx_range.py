import docx

doc = docx.Document("/home/cubeai/Videos/Dinesh Export Textile /Dinesh_Textile(may 26)/WEAVING_YARN_CALCULATION_GUIDE.docx")

for idx in range(70, min(100, len(doc.paragraphs))):
    print(f"P{idx}: {doc.paragraphs[idx].text}")
