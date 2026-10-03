import os
import sys
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def build_workflow_document(output_path):
    print(f"Creating document: {output_path}")
    doc = docx.Document()

    # Set page margins to 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Styles & Colors
    C_PRIMARY = RGBColor(30, 58, 138)      # Deep Navy #1E3A8A
    C_SECONDARY = RGBColor(13, 148, 136)   # Teal #0D9488
    C_ACCENT = RGBColor(67, 56, 202)       # Indigo #4338CA
    C_DARK = RGBColor(31, 41, 55)          # Dark Slate #1F2937
    C_MUTED = RGBColor(107, 114, 128)      # Gray #6B7280
    HEX_PRIMARY = "1E3A8A"
    HEX_SECONDARY = "0D9488"
    HEX_LIGHT_BG = "F3F4F6"
    HEX_ALT_ROW = "F9FAFB"

    def set_cell_background(cell, fill_hex):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m_name, m_val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m_name}')
            node.set(qn('w:w'), str(m_val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'  <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:left w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'  <w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def format_table_header(row, col_widths, titles):
        for idx, cell in enumerate(row.cells):
            cell.width = col_widths[idx]
            set_cell_background(cell, HEX_PRIMARY)
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(titles[idx])
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(255, 255, 255)

    def format_table_row(row, col_widths, values, is_even=False):
        fill_color = HEX_ALT_ROW if is_even else "FFFFFF"
        for idx, cell in enumerate(row.cells):
            cell.width = col_widths[idx]
            set_cell_background(cell, fill_color)
            set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(str(values[idx]))
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.color.rgb = C_DARK

    def add_heading_1(text):
        h = doc.add_heading(level=1)
        h.paragraph_format.space_before = Pt(20)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(15)
        run.bold = True
        run.font.color.rgb = C_PRIMARY
        return h

    def add_heading_2(text):
        h = doc.add_heading(level=2)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(12.5)
        run.bold = True
        run.font.color.rgb = C_SECONDARY
        return h

    def add_heading_3(text):
        h = doc.add_heading(level=3)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10.5)
        run.bold = True
        run.font.color.rgb = C_ACCENT
        return h

    def add_paragraph(text, bold_prefix=None, italic=False, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(10)
            r_pre.bold = True
            r_pre.font.color.rgb = C_DARK
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.italic = italic
        r.font.color.rgb = C_DARK
        return p

    def add_bullet_point(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Arial"
            r_pre.font.size = Pt(9.5)
            r_pre.bold = True
            r_pre.font.color.rgb = C_DARK
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(9.5)
        r.font.color.rgb = C_DARK
        return p

    def add_callout(text, title="WORKFLOW HIGHLIGHT", fill_hex="F0FDF4", border_color="16A34A"):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, fill_hex)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=160)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
            f'  <w:top w:val="none"/>'
            f'  <w:right w:val="none"/>'
            f'  <w:bottom w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        r_title = p.add_run(f"[{title}]  ")
        r_title.bold = True
        r_title.font.name = "Arial"
        r_title.font.size = Pt(9.5)
        r_title.font.color.rgb = RGBColor(22, 163, 74) if border_color == "16A34A" else RGBColor(37, 99, 235)
        
        r_text = p.add_run(text)
        r_text.font.name = "Arial"
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = C_DARK
        
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_code_block(text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.cell(0, 0)
        cell.width = Inches(6.5)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'  <w:left w:val="single" w:sz="16" w:space="0" w:color="0284C7"/>'
            f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(text)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # -------------------------------------------------------------
    # 1. COVER PAGE
    # -------------------------------------------------------------
    cov_p1 = doc.add_paragraph()
    cov_p1.paragraph_format.space_before = Pt(60)
    cov_p1.paragraph_format.space_after = Pt(12)
    cov_r1 = cov_p1.add_run("HANDLOOM ENTERPRISE RESOURCE PLANNING")
    cov_r1.font.name = "Arial"
    cov_r1.font.size = Pt(24)
    cov_r1.bold = True
    cov_r1.font.color.rgb = C_PRIMARY

    cov_p2 = doc.add_paragraph()
    cov_p2.paragraph_format.space_before = Pt(0)
    cov_p2.paragraph_format.space_after = Pt(24)
    cov_r2 = cov_p2.add_run("Comprehensive End-to-End System Workflow, Functional Architecture, and Operational Manual")
    cov_r2.font.name = "Arial"
    cov_r2.font.size = Pt(13)
    cov_r2.font.color.rgb = C_SECONDARY

    cov_bar = doc.add_paragraph()
    cov_bar.paragraph_format.space_before = Pt(0)
    cov_bar.paragraph_format.space_after = Pt(36)
    r_bar = cov_bar.add_run("―" * 45)
    r_bar.font.color.rgb = C_SECONDARY
    r_bar.bold = True

    # Metadata table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_borders(meta_table, color="CBD5E1", sz="4")
    m_widths = [Inches(2.0), Inches(4.5)]
    m_data = [
        ("Project Name", "Handloom ERP (Textile Manufacturing & Financial System)"),
        ("Document Version", "1.0.0 (Production Release)"),
        ("Target Industry", "Textile Manufacturing, Weaving, Dyeing, Export & Trading"),
        ("System Architecture", "FastAPI (Async Python 3.13) + React 18 (Vite) + PostgreSQL / SQLite"),
        ("Integrated Modules", "Design AI, Merchandising, PPC, Multi-Stage Job Work, Stores, Fleet, CubeBook Finance"),
        ("Publication Date", datetime.now().strftime("%B %d, %Y"))
    ]
    for idx, (k, v) in enumerate(m_data):
        row = meta_table.rows[idx]
        for c_idx, cell in enumerate(row.cells):
            cell.width = m_widths[c_idx]
            set_cell_background(cell, "F8FAFC" if c_idx == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(k if c_idx == 0 else v)
            run.font.name = "Arial"
            run.font.size = Pt(9.5)
            if c_idx == 0:
                run.bold = True
                run.font.color.rgb = C_PRIMARY
            else:
                run.font.color.rgb = C_DARK

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. EXECUTIVE SUMMARY & BUSINESS CONTEXT
    # -------------------------------------------------------------
    add_heading_1("1. Executive Summary & Business Context")
    
    add_paragraph(
        "Handloom ERP is an enterprise-grade resource planning system specifically designed for the complex, "
        "multi-tiered textile and handloom manufacturing industry. Textile manufacturing differs fundamentally from standard discrete "
        "manufacturing due to continuous material transformations: cotton fibers become single yarns, yarns undergo twisting and doubling, "
        "grey yarns are chemically dyed, warp beams are sized and warped, weft yarns are interwoven on handlooms or automatic looms into "
        "greige fabrics, and greige fabrics pass through multiple wet processing stages (dyeing, printing, calendering, sanforizing) before "
        "reaching finished fabric status.",
        bold_prefix="Business Context: "
    )
    
    add_paragraph(
        "Traditional textile mills and handloom societies operate via fragmented manual registers and Excel spreadsheets. This leads to critical "
        "operational leakage: unaccounted yarn wastage at job work facilities, lack of visibility into loom downtime and weaving progress, "
        "discrepancies between greige meterage and finished fabric output, delays in GST/E-Way bill generation, and complete disconnection "
        "between factory floor movements and financial accounting ledgers."
    )

    add_paragraph(
        "Handloom ERP bridges this divide by delivering a unified digital platform. It orchestrates every phase from initial Buyer Order "
        "booking, CAD design pattern extraction via AI vision, yarn consumption calculation, yarn purchase orders, multi-stage job work delivery challans "
        "(Twisting, Yarn Dyeing, Warping & Sizing, Weaving, Fabric Dyeing, Finishing), on-table 4-point quality inspection, warehouse rack tracking, "
        "packing slips, sales invoicing, fleet dispatch, to real-time double-entry bookkeeping in CubeBook (a native Tally-compatible finance engine).",
        bold_prefix="Core Mission: "
    )

    add_callout(
        "Zero Material Leakage: Handloom ERP maintains strict material accountability across every external job worker and internal machine shed. "
        "Every gram of yarn issued for dyeing, warping, or weaving must reconcile with receipts, process shrinkage, and reported wastage.",
        title="CORE VALUE PROPOSITION",
        fill_hex="EFF6FF",
        border_color="2563EB"
    )

    # -------------------------------------------------------------
    # 3. SYSTEM ARCHITECTURE & TECHNICAL STACK
    # -------------------------------------------------------------
    add_heading_1("2. System Architecture & Technical Infrastructure")
    
    add_paragraph(
        "The application is architected around high-performance modern web standards, featuring an asynchronous backend API, a responsive "
        "component-driven frontend, and a dual-database design ensuring both high-throughput production recording and robust financial integrity."
    )

    add_heading_2("2.1 Technical Stack Overview")
    
    tech_table = doc.add_table(rows=1, cols=3)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tech_table)
    t_widths = [Inches(1.5), Inches(2.2), Inches(2.8)]
    format_table_header(tech_table.rows[0], t_widths, ["Layer", "Technology / Framework", "Key Responsibilities"])
    
    tech_specs = [
        ("Frontend Client", "React 18, Vite, Tailwind CSS, Lucide Icons", "Responsive SPA, dynamic master forms, live PPC boards, print previews, RBAC guards."),
        ("Backend Services", "FastAPI (Python 3.13), Uvicorn", "High-throughput async REST API, Pydantic v2 validation, HTTP audit logging middleware."),
        ("ORM & Data Layer", "SQLAlchemy 2.0 (Async), Alembic", "Async PostgreSQL driver (asyncpg) for core ERP; SQLite (aiosqlite) for CubeBook Finance."),
        ("AI Vision Module", "OpenCV, Scikit-Learn (KMeans), SciPy, PIL", "Computer vision fabric pattern recognition, dominant color clustering, 2D FFT weave structure analysis."),
        ("Document Generation", "WeasyPrint, XHTML2PDF, OpenPyXL", "A4 PDF invoices, Gate Passes, Delivery Challans, Barcode packing slips, and Excel MIS reports."),
        ("Security & Auth", "JWT (python-jose), Passlib (Bcrypt)", "Role-Based Access Control (RBAC), token blacklisting, password salting, granular module permissions.")
    ]
    for idx, row_data in enumerate(tech_specs):
        row = tech_table.add_row()
        format_table_row(row, t_widths, row_data, is_even=(idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    add_heading_2("2.2 Dual-Database Architecture & Real-Time Sync Engine")
    add_paragraph(
        "To balance the high velocity of shop-floor updates (daily pick counts, gate entries, inspection grades) with the immutability demanded by financial "
        "accounting, Handloom ERP implements a synchronized dual-database architecture:"
    )
    add_bullet_point("Stores all operational entities: Party Master, Buyer Orders, Yarn Purchase Orders, Job Work Challans, Grey Inspection, Warehouse Stocks, Fleet, and HR.", bold_prefix="Primary PostgreSQL Database: ")
    add_bullet_point("Embedded financial engine modeled after Tally ERP standards, holding Companies, Ledger Groups, Ledgers, Vouchers (Payment, Receipt, Journal, Contra, Sales, Purchase), and Financial Statements.", bold_prefix="CubeBook SQLite Database: ")
    add_bullet_point("An event-driven synchronization service (`sync_realtime_data.py`) listens to operational milestones. When a Sales Invoice is finalized or a Yarn GRN is accepted, it automatically constructs balanced double-entry accounting vouchers in CubeBook without requiring manual accountant intervention.", bold_prefix="Real-Time Sync Engine: ")

    add_code_block(
        "+-----------------------------------------------------------------------------------+\n"
        "|                             CLIENT LAYER: React 18 SPA                           |\n"
        "|     [ Merchandising ] [ PPC Engine ] [ Job Work ] [ Inventory ] [ CubeBook ]      |\n"
        "+-----------------------------------------+-----------------------------------------+\n"
        "                                          | HTTP / REST (JWT Auth)\n"
        "                                          v\n"
        "+-----------------------------------------------------------------------------------+\n"
        "|                          API GATEWAY: FastAPI (Python 3.13)                       |\n"
        "|  +--------------------+  +----------------------+  +---------------------------+  |\n"
        "|  | RBAC & Auth Guard  |  | AI Vision Analytics  |  | Audit Trail Middleware    |  |\n"
        "|  +--------------------+  +----------------------+  +---------------------------+  |\n"
        "+-------------------+-------------------------------------+-------------------------+\n"
        "                    |                                     |\n"
        "                    v                                     v\n"
        "+------------------------------------+   +------------------------------------------+\n"
        "|     CORE ERP ENGINE (PostgreSQL)   |   |        FINANCE SUB-APP (CubeBook SQLite) |\n"
        "| - Buyer Orders & Merchandising     |   | - Chart of Accounts & Ledgers            |\n"
        "| - Yarn Purchase & Job Work Challans|   | - Double-Entry Financial Vouchers        |\n"
        "| - 4-Point Quality Inspection       |<--+ - Real-Time Automated Sync Pipeline      |\n"
        "| - Warehouse, Fleet & HR Modules    |   | - GST (GSTR-1/3B) & Balance Sheet        |\n"
        "+------------------------------------+   +------------------------------------------+"
    )

    # -------------------------------------------------------------
    # 4. MASTER DATA MANAGEMENT
    # -------------------------------------------------------------
    add_heading_1("3. Master Data Management (The Operational Foundation)")
    add_paragraph(
        "Before any operational transactions can commence, the system relies on comprehensive Master Data Management (MDM) to configure enterprise entities, "
        "manufacturing parameters, quality tolerances, and business partner relations."
    )

    add_heading_2("3.1 Party Master")
    add_paragraph(
        "The Party Master serves as the single source of truth for all external stakeholders interacting with the enterprise. Each party is categorized "
        "by role, tax profile, and credit terms:"
    )
    add_bullet_point("International and domestic apparel brands, retail chains, and garment buying houses.", bold_prefix="Buyers / Customers: ")
    add_bullet_point("Spinning mills and yarn traders supplying Cotton, Silk, Linen, Viscose, and Blended yarns.", bold_prefix="Yarn Suppliers: ")
    add_bullet_point("Third-party processing units providing twisting, cone dyeing, warp sizing, handloom/powerloom weaving, stenter finishing, and rotary printing.", bold_prefix="Job Workers: ")
    add_bullet_point("Logistics companies providing trucks, containers, and local transport.", bold_prefix="Transporters: ")
    add_bullet_point("Suppliers of spare parts, dyes, sizing chemicals, stationery, and packing boxes.", bold_prefix="General Vendors: ")

    add_heading_2("3.2 38+ Configurable Sub-Masters")
    add_paragraph(
        "To maintain data cleanliness and prevent free-text discrepancies, all operational parameters are governed by generic sub-master entities:"
    )
    
    sub_table = doc.add_table(rows=1, cols=3)
    sub_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sub_table)
    s_widths = [Inches(1.8), Inches(2.3), Inches(2.4)]
    format_table_header(sub_table.rows[0], s_widths, ["Domain", "Master Entities", "Business Purpose"])
    
    sub_data = [
        ("Yarn & Raw Material", "Yarn Count Master, Yarn Ply, Fiber Type, Filament Type", "Standardizes counts (e.g., 2/40s Cotton, 60s Linen) and ply construction."),
        ("Weaving & Loom", "Loom Type (Handloom, Powerloom, Rapier), Reed Count, Pick Master", "Calculates machine capability, reed width capacity, and pick density."),
        ("Design & Fabric", "Weave Structure (Plain, Twill, Satin, Jacquard), Fabric Category, Color Master", "Determines fabric construction parameters and visual properties."),
        ("Commercial & Export", "Currencies (INR, USD, EUR, GBP), Payment Terms, Delivery Terms (FOB, CIF)", "Governs international export pricing and credit duration."),
        ("Organization & Staff", "Department Master (19 depts), Designation Master (20 roles), Shift Master", "Enforces organizational hierarchy, reporting lines, and HR shift rotations.")
    ]
    for idx, row_data in enumerate(sub_data):
        row = sub_table.add_row()
        format_table_row(row, s_widths, row_data, is_even=(idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 5. END-TO-END OPERATIONAL WORKFLOW (STEP-BY-STEP)
    # -------------------------------------------------------------
    add_heading_1("4. End-to-End Operational Lifecycle Workflow")
    add_paragraph(
        "The operational lifecycle of Handloom ERP spans 13 cohesive phases that guide raw yarn through pre-production planning, multi-facility "
        "manufacturing, rigorous quality checking, customer dispatch, and financial realization."
    )

    # PHASE 1
    add_heading_2("Phase 1: Design Management & Pre-Production Engineering")
    add_paragraph(
        "Every fabric manufactured originates in the Design Management module. Before committing capital to yarn purchases, the technical construction "
        "and unit economics must be locked."
    )
    add_bullet_point("Designers upload high-resolution fabric scans or CAD swatches. The system records Reed Count, Picks Per Inch (PPI), Ends Per Inch (EPI), Total Ends, Reed Space, and Finished Width.", bold_prefix="Step 1.1 - Design Entry: ")
    add_bullet_point("Integrated computer vision algorithms analyze the uploaded image. A K-Means clustering algorithm extracts the dominant color palette with percentage ratios. 2D Fast Fourier Transform (FFT) detects periodicity in the warp/weft grid to verify weave structure (plain, twill, satin).", bold_prefix="Step 1.2 - AI Vision Pattern Analysis: ")
    add_bullet_point("Calculates exact material requirements using textile engineering formulas: Warp weight in Kgs = (Total Ends × Tape Length in Meters) / (1693.3 × Yarn Count × (1 - Crimp%)). Weft weight in Kgs = (Picks/Inch × Reed Space in Inches × Fabric Length) / (1693.3 × Weft Count).", bold_prefix="Step 1.3 - Weaving Calculator: ")
    add_bullet_point("Compiles a multi-layered cost breakdown: Raw yarn cost + Twisting charges + Yarn Dyeing charges + Warping & Sizing + Weaving charges per pick/meter + Wet processing & finishing + Packaging + Administrative overheads. Incorporates target profit margins to compute final sales quotation rates.", bold_prefix="Step 1.4 - Costing Sheet & Masters: ")

    # PHASE 2
    add_heading_2("Phase 2: Merchandising & Buyer Order Lifecycle")
    add_paragraph(
        "Once the design and costing are accepted by the customer, the order enters the merchandising pipeline."
    )
    add_bullet_point("Merchandiser creates the formal Buyer Order capturing Buyer Code, Purchase Order Number, Order Date, Currency, Incoterms (FOB, CIF, Ex-Factory), and detailed itemized lines with delivery deadlines.", bold_prefix="Step 2.1 - Buyer Order Booking: ")
    add_bullet_point("Large orders are split into shipment tranches and lot allocations. The system plans staggered production targets across weeks to meet multiple delivery dates.", bold_prefix="Step 2.2 - Order Scheduling: ")
    add_bullet_point("A formal Proforma Invoice (PI) is generated from the order details. The system triggers an approval workflow requiring approval from the Merchandising Manager and Credit Control before production release.", bold_prefix="Step 2.3 - Proforma Invoice (PI) & Approval: ")
    add_bullet_point("If the customer requests modifications to colorways, meterage, or shipping dates, a versioned amendment is created, retaining full audit history of prior revisions.", bold_prefix="Step 2.4 - Order Amendment Tracking: ")
    add_bullet_point("Incidental costs incurred specifically for the order (courier of sample swatches, specialized lab test certifications, buyer inspection agency fees) are logged directly against the order to compute net order profitability.", bold_prefix="Step 2.5 - Order Expenses: ")

    # PHASE 3
    add_heading_2("Phase 3: Production Planning & Control (PPC)")
    add_paragraph(
        "The PPC module orchestrates machine allocation, prevents bottlenecks, and forecasts delivery fulfillment."
    )
    add_bullet_point("Maintains an inventory of all weaving sheds, handloom clusters, and automatic looms with reed width limits, current speed (PPM), and standard efficiency benchmarks.", bold_prefix="Step 3.1 - Loom Master Setup: ")
    add_bullet_point("The PPC planning algorithm evaluates outstanding order schedules against idle loom capacity. Orders are assigned to specific looms based on warp width, beam length, and weaver specialization.", bold_prefix="Step 3.2 - Loom Allocation: ")
    add_bullet_point("Provides real-time visibility into active lots, current shift output, loom downtime, and efficiency variances across factory sheds.", bold_prefix="Step 3.3 - Live PPC Dashboard: ")
    add_bullet_point("Computes dynamic completion forecasts based on historical weaver output, active picks per minute, and remaining beam meterage. Automatically alerts supervisors if an order is projected to miss its dispatch milestone.", bold_prefix="Step 3.4 - ETA Calculation Engine: ")
    add_bullet_point("Operators and supervisors log loom stoppages (warp breakages, mechanical failure, power outage, shortage of weft yarn cones). Stoppages feed into maintenance alerts and downtime reports.", bold_prefix="Step 3.5 - Breakdown & Problem Entry: ")

    # PHASE 4
    add_heading_2("Phase 4: Raw Material Procurement & Inward")
    add_paragraph(
        "Raw materials (grey yarn cones, dyed yarn, and external grey fabrics) are procured systematically."
    )
    add_bullet_point("Purchase Manager generates Yarn Purchase Orders (YPO) linking directly to required Buyer Orders. Specifies yarn count, fiber blend, mill brand, delivery terms, and agreed per-Kg rates.", bold_prefix="Step 4.1 - Yarn Purchase Order (YPO): ")
    add_bullet_point("When yarn trucks arrive at the factory gate, a Gate Inward Pass is generated. The Stores team conducts gross, tare, and net weighment, counts bags/cartons, verifies cone counts, and checks mill test certificates.", bold_prefix="Step 4.2 - Yarn Inward & Goods Receipt Note (GRN): ")
    add_bullet_point("Accepted yarn lots are assigned specific Rack and Bin locations in the warehouse. The physical inventory balance updates immediately.", bold_prefix="Step 4.3 - Rack Master & Bin Allocation: ")
    add_bullet_point("The system tracks raw material inventory across warehouses. Inward data triggers an automated purchase entry in the CubeBook finance module.", bold_prefix="Step 4.4 - Inventory Ledger Sync: ")

    # PHASE 5
    add_heading_2("Phase 5: Multi-Stage Job Work & Manufacturing Pipeline")
    add_paragraph(
        "Handloom textile manufacturing involves a multi-step conversion process, often distributed across specialized job work facilities. "
        "Handloom ERP maintains strict material accounting at each phase:"
    )
    
    jw_table = doc.add_table(rows=1, cols=4)
    jw_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(jw_table)
    jw_widths = [Inches(1.2), Inches(1.8), Inches(1.8), Inches(1.7)]
    format_table_header(jw_table.rows[0], jw_widths, ["Stage", "Process / PO", "Outward Delivery", "Inward Receipt & Control"])
    
    jw_steps = [
        ("Stage 5.1", "Twisting & Doubling PO", "Single ply grey yarn issued to twister via Delivery Challan (DC).", "2-ply or multi-ply twisted yarn received; twisting loss reconciled."),
        ("Stage 5.2", "Yarn Dyeing PO", "Grey yarn delivered to cone/hank dyeing master with approved Lab Dip shade recipe.", "Dyed yarn cones received; shade matched; dyeing weight gain/loss calculated."),
        ("Stage 5.3", "Warping & Sizing PO", "Dyed/grey warp yarn cones issued to warping and sizing unit.", "Warp beams received with beam number, total ends count, and beam length in meters."),
        ("Stage 5.4", "Weaving PO", "Warp beams and weft yarn cones delivered to handloom/powerloom weavers.", "Woven greige cloth received (Weaving Receipt / Cloth Inward); pick density verified."),
        ("Stage 5.5", "Fabric Dyeing PO", "Greige cloth issued to fabric processing house.", "Solid dyed fabric received; length shrinkage and width checked."),
        ("Stage 5.6", "Printing PO", "Dyed/bleached fabric delivered to screen or rotary printing unit.", "Printed fabric received; print registration and color fastness inspected."),
        ("Stage 5.7", "Finishing PO", "Cloth sent for Stenter, Sanforizing, Mercerizing, or Calendering.", "Finished fabric received in ready-to-cut condition; final shrinkage fixed.")
    ]
    for idx, row_data in enumerate(jw_steps):
        row = jw_table.add_row()
        format_table_row(row, jw_widths, row_data, is_even=(idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    add_paragraph(
        "Job workers submit bills against their agreed rate contracts. The system reconciles the billed quantity with recorded inward weights "
        "and Delivery Challan references, deducts applicable Tax Deducted at Source (TDS), and passes the approved liability to Accounts.",
        bold_prefix="Step 5.8 - Job Work Bill Passing: "
    )

    # PHASE 6
    add_heading_2("Phase 6: Quality Control, Inspection & Fabric Grading")
    add_paragraph(
        "Quality control safeguards against buyer rejections and export penalties. Handloom ERP implements an automated 4-Point System:"
    )
    add_bullet_point("Inspectors mount fabric rolls on illuminated inspection tables. The inspector enters defects per 100 square meters in real time: slubs, broken picks, float, reed marks, oil spots, and color fly.", bold_prefix="Step 6.1 - On-Table Checking (Grey Inspection): ")
    add_bullet_point("Faults are scored 1 to 4 penalty points based on size. The system calculates total points per 100 square meters. Rolls exceeding threshold tolerances are automatically downgraded.", bold_prefix="Step 6.2 - Automated 4-Point Grading: ")
    add_bullet_point("Rolls are classified into Fresh (A-Grade export quality), Seconds (B-Grade discounted market sale), Rejection (C-Grade unsaleable), or Cut-Pieces (salvageable remnants).", bold_prefix="Step 6.3 - Classification: ")
    add_bullet_point("Post-finishing inspection verifies finished width (inches), GSM (Grams per Square Meter), warp/weft shrinkage %, and shade consistency under standard D65 light sources.", bold_prefix="Step 6.4 - Finished Fabric Inspection: ")
    add_bullet_point("Each inspected roll receives a unique piece barcode number storing meterage, gross weight, tare weight, grade, and inspector ID.", bold_prefix="Step 6.5 - Barcode Identification: ")

    # PHASE 7
    add_heading_2("Phase 7: Warehouse, Godowns & Inventory Management (DDD Module)")
    add_paragraph(
        "The warehouse architecture follows Domain-Driven Design (DDD) principles to maintain absolute material control:"
    )
    add_bullet_point("Hierarchical organization of multiple locations: Main Yarn Godown, Sizing Staging Shed, Weaving WIP Storage, Greige Mending Godown, Finished Fabric Warehouse, and Export Packing Hub.", bold_prefix="Step 7.1 - Multi-Godown Hierarchy: ")
    add_bullet_point("Dedicated physical unloading zone where incoming cargo is quarantined, verified against vendor challans, and barcoded before put-away.", bold_prefix="Step 7.2 - Goods Receipt Dock: ")
    add_bullet_point("System recommends optimal rack, tier, and bin assignments based on fabric category, turnover frequency, and batch shelf-life.", bold_prefix="Step 7.3 - Put-Away Management: ")
    add_bullet_point("Warehouse personnel receive digitized pick lists for order packing or job work issuance, enforcing First-In, First-Out (FIFO) material movement.", bold_prefix="Step 7.4 - Pick List Generation: ")
    add_bullet_point("Dedicated registers for Raw Materials, WIP across external units, Finished Goods, Consumables, Machine Spares, and Surplus/Dead Stock with periodic physical stock reconciliation.", bold_prefix="Step 7.5 - Stock Ledgers & Audit: ")

    # PHASE 8
    add_heading_2("Phase 8: Packing, Dispatch & Commercial Invoicing")
    add_paragraph(
        "When production and quality inspection are concluded, goods are prepared for commercial shipment:"
    )
    add_bullet_point("Pieces are rolled, poly-wrapped, and packed into master cartons or hessian-wrapped bales. The packing slip records box numbers, piece IDs, net weight, gross weight, and meterage.", bold_prefix="Step 8.1 - Packing Slip Creation: ")
    add_bullet_point("Quality and Accounts teams authorize a Goods Release Advice (GRA) confirming the customer has met payment/credit terms and quality checks are cleared.", bold_prefix="Step 8.2 - Goods Release Advice (GRA): ")
    add_bullet_point("The dispatch team assigns a transporter, vehicle number, and shipping route, finalizing the container loading sheet.", bold_prefix="Step 8.3 - Despatch Planning: ")
    add_bullet_point("The system generates a GST-compliant Tax Invoice featuring Buyer GSTIN, HSN Code, Taxable Value, CGST + SGST (intra-state) or IGST (inter-state), TCS, freight charges, and bank remittance instructions.", bold_prefix="Step 8.4 - Sales Invoice Generation: ")
    add_bullet_point("Automated generation of E-Way Bills complying with national tax portal schemas, capturing Distance, Transporter ID, and Part A/Part B vehicle details.", bold_prefix="Step 8.5 - E-Way Bill Integration: ")

    # PHASE 9
    add_heading_2("Phase 9: Fleet, Transportation & Logistics Management")
    add_paragraph(
        "For enterprises operating proprietary delivery vehicles or contracted truck fleets, the Fleet module manages logistics operations:"
    )
    add_bullet_point("Tracks company trucks, tempos, and contracted vehicles with tonnage capacity, fuel type, and registration numbers.", bold_prefix="Step 9.1 - Vehicle & Driver Master: ")
    add_bullet_point("Plans multi-stop delivery runs connecting spinning mills, dyeing plants, weaving sheds, and customer ports.", bold_prefix="Step 9.2 - Trip Planning & Execution: ")
    add_bullet_point("Drivers log fuel slips, odometer readings, and toll receipts. The system computes KM per liter and cost per metric ton-KM.", bold_prefix="Step 9.3 - Fuel & Mileage Analysis: ")
    add_bullet_point("Maintains service schedules, preventive maintenance records, and unexpected breakdown logs.", bold_prefix="Step 9.4 - Maintenance Logs: ")
    add_bullet_point("Automated alert engine monitors expiring vehicle insurance, fitness certificates, road tax, and driver commercial licenses.", bold_prefix="Step 9.5 - Document Expiry Alerts: ")

    # PHASE 10
    add_heading_2("Phase 10: Stores, Spares & Consumables Management")
    add_paragraph(
        "Non-yarn inventory (shuttles, heald wires, reeds, sizing starch, dyes, printing paste, stationery, packing cartons) is governed through Stores:"
    )
    add_bullet_point("Departments (Weaving, Dyeing, Maintenance, Admin) submit digitized material requisitions for consumables.", bold_prefix="Step 10.1 - Material Requisition: ")
    add_bullet_point("Store Manager consolidates requests into a Purchase Requisition, solicits vendor quotations, and issues Purchase Orders upon approval.", bold_prefix="Step 10.2 - Requisition to PO: ")
    add_bullet_point("Stock Inward GRN is performed at the store, followed by barcode tagging.", bold_prefix="Step 10.3 - Store GRN Inward: ")
    add_bullet_point("Items are issued directly to cost centers (e.g., Weaving Shed 2, Stenter Line 1). Unused materials are returned or transferred.", bold_prefix="Step 10.4 - Issue & Consumption Tracking: ")
    add_bullet_point("Controls swatch card distribution to buyers and returnable delivery challans (RDC) for tools and samples.", bold_prefix="Step 10.5 - Swatch Cards & Returnable DC: ")

    # PHASE 11
    add_heading_2("Phase 11: Human Resources & Payroll Management (HR)")
    add_paragraph(
        "Manages workforce operations across management staff, master weavers, technicians, and daily wage workers:"
    )
    add_bullet_point("Stores employee profiles, biometric IDs, bank accounts, statutory PF/ESI numbers, and salary structures.", bold_prefix="Step 11.1 - Employee Master: ")
    add_bullet_point("Integrates shift rosters, daily attendance, overtime hours, and leave applications with automated approval workflows.", bold_prefix="Step 11.2 - Attendance & Shift Tracking: ")
    add_bullet_point("Computes monthly salaries with earnings (Basic, HRA, Conveyance, Allowances) and deductions (PF, ESI, Professional Tax, TDS, Loan recovery), generating PDF pay slips.", bold_prefix="Step 11.3 - Payroll Processing: ")
    add_bullet_point("Tracks staff salary advances, festival loans, travel approvals, and expense reimbursement claims.", bold_prefix="Step 11.4 - Employee Loans & Claims: ")

    # PHASE 12
    add_heading_2("Phase 12: Integrated Financial Accounting (CubeBook)")
    add_paragraph(
        "CubeBook is the native accounting sub-application embedded within Handloom ERP. It operates as a full-featured, Tally-compatible double-entry "
        "general ledger system:"
    )
    add_bullet_point("Multi-company support with customizable financial year boundaries (e.g., April 1 to March 31).", bold_prefix="Step 12.1 - Company & Financial Years: ")
    add_bullet_point("Organized into Primary Groups (Current Assets, Fixed Assets, Current Liabilities, Direct Expenses, Sales Accounts, Purchase Accounts) with hierarchical sub-ledgers.", bold_prefix="Step 12.2 - Chart of Accounts: ")
    add_bullet_point("Full dual-entry support for Payment, Receipt, Contra, Journal, Sales, Purchase, Credit Note, and Debit Note vouchers.", bold_prefix="Step 12.3 - Comprehensive Voucher Engine: ")
    add_bullet_point("Every finalized operational transaction automatically posts balanced accounting entries. A Sales Invoice debits Sundry Debtors and credits Sales Account + CGST/SGST/IGST Payable. A Yarn GRN debits Yarn Purchase Account and credits Sundry Creditors.", bold_prefix="Step 12.4 - Real-Time Automated Sync: ")
    add_bullet_point("Import bank statements, log cheque payments, and perform bank reconciliation statements (BRS).", bold_prefix="Step 12.5 - Banking & Reconciliation: ")
    add_bullet_point("Generates GSTR-1 (Outward Supplies), GSTR-3B monthly summary, and Input Tax Credit (ITC) reconciliation.", bold_prefix="Step 12.6 - Statutory GST Reporting: ")
    add_bullet_point("Real-time generation of Day Book, Trial Balance, Profit & Loss Statement, Balance Sheet, Debtor/Creditor Ageing Analysis, and Financial Ratios.", bold_prefix="Step 12.7 - Financial Statements: ")

    # PHASE 13
    add_heading_2("Phase 13: System Administration, Security & Audit Governance")
    add_paragraph(
        "Enterprise security ensures data privacy, access control, and complete traceability across all users:"
    )
    add_bullet_point("Fine-grained Role-Based Access Control (RBAC) assigning permissions (View, Create, Edit, Delete, Approve) on a per-module basis. Pre-seeded roles include Super Admin, Admin, Merchandiser, Weaving Supervisor, Quality Inspector, Storekeeper, Dispatch Officer, and Accountant.", bold_prefix="Step 13.1 - Role-Based Access Control: ")
    add_bullet_point("An asynchronous HTTP middleware intercepts all mutating requests (POST, PUT, DELETE). It logs the authenticated user ID, client IP, timestamp, module name, HTTP action, and payload summary to the immutable Log Report table.", bold_prefix="Step 13.2 - Automated Audit Trail Middleware: ")
    add_bullet_point("Alerts users to pending approvals (PO Approval, GRA Approval, PI Sign-off) and system triggers (loom breakdown, document expiry).", bold_prefix="Step 13.3 - Real-Time Notifications Hub: ")

    # -------------------------------------------------------------
    # 6. VISUAL PROCESS FLOWCHARTS & LIFECYCLE MAPS
    # -------------------------------------------------------------
    add_heading_1("5. Visual Process Maps & Lifecycle Diagrams")
    add_paragraph(
        "The following process maps illustrate the core business lifecycles and material flows running through Handloom ERP."
    )

    add_heading_2("5.1 Complete Textile Manufacturing Value Stream")
    add_code_block(
        "[ 1. INQUIRY & CAD DESIGN ] ---> [ 2. AI PATTERN ANALYSIS ] ---> [ 3. WEAVING CALCULATOR ]\n"
        "                                                                         |\n"
        "                                                                         v\n"
        "[ 6. BUYER ORDER SCHEDULE ] <--- [ 5. BUYER ORDER BOOKING ] <--- [ 4. COSTING SHEET ]\n"
        "             |\n"
        "             v\n"
        "[ 7. PPC LOOM PLANNING ] ------> [ 8. YARN PURCHASE ORDER ] ----> [ 9. GATE INWARD & GRN ]\n"
        "                                                                         |\n"
        "                                                                         v\n"
        "[ 12. WARP BEAM RECEIPT ] <--- [ 11. WARPING & SIZING PO ] <--- [ 10. YARN DYEING PO ]\n"
        "             |\n"
        "             v\n"
        "[ 13. WEAVING PO & SHED ] -----> [ 14. GREIGE CLOTH INWARD ] ---> [ 15. 4-POINT INSPECTION ]\n"
        "                                                                         |\n"
        "                                                                         v\n"
        "[ 18. BARCODE PACKING ] <------- [ 17. FINISHING & STENTER ] <--- [ 16. FABRIC DYEING PO ]\n"
        "             |\n"
        "             v\n"
        "[ 19. GOODS RELEASE (GRA) ] ---> [ 20. SALES INVOICE & E-WAY ] -> [ 21. CUBEBOOK LEDGERS ]"
    )

    add_heading_2("5.2 Order-to-Cash (O2C) Commercial Flow")
    add_code_block(
        "Buyer Inquiry -> Costing Sheet -> Buyer Order Booking -> Proforma Invoice (PI) Approval\n"
        "  -> PPC Production Execution -> Quality Clearance -> Packing Slip -> Goods Release Advice (GRA)\n"
        "  -> Sales Invoice & E-Way Bill -> Transporter Dispatch -> Financial Sync (Debit Customer / Credit Sales)\n"
        "  -> Payment Receipt Voucher -> Bank Statement Reconciliation."
    )

    add_heading_2("5.3 Procure-to-Pay (P2P) Raw Material Flow")
    add_code_block(
        "Production Material Requirement -> Purchase Requisition -> Yarn Purchase Order (YPO)\n"
        "  -> Gate Inward Verification -> GRN & Weighment -> Warehouse Rack Allocation\n"
        "  -> Financial Sync (Debit Yarn Inventory / Credit Supplier) -> Vendor Bill Passing\n"
        "  -> Payment Voucher Disbursed."
    )

    add_heading_2("5.4 Job Work Material Balancing Loop")
    add_code_block(
        "Issued Raw Yarn (Kg)  ===========================================> Job Worker Facility\n"
        "                              |                                          |\n"
        "                              v                                          v\n"
        "Received Processed Yarn (Kg) + Allowable Loss (Kg) + Wastage Scrap (Kg) = Total Accountability"
    )

    # -------------------------------------------------------------
    # 7. OPERATIONAL RACI MATRIX
    # -------------------------------------------------------------
    add_heading_1("6. Operational RACI Matrix (Roles & Responsibilities)")
    add_paragraph(
        "The RACI Matrix defines operational governance across key roles: Responsible (R), Accountable (A), Consulted (C), and Informed (I)."
    )
    
    raci_table = doc.add_table(rows=1, cols=7)
    raci_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(raci_table)
    r_widths = [Inches(1.8), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.8), Inches(0.8)]
    format_table_header(raci_table.rows[0], r_widths, ["Workflow Stage", "Merch.", "PPC Mgr", "Prod. Sup", "QC Insp.", "Stores", "Accounts"])
    
    raci_data = [
        ("Design Entry & Costing", "R / A", "C", "I", "C", "I", "C"),
        ("Buyer Order & PI Booking", "R / A", "C", "I", "I", "I", "C"),
        ("Order Approval (PI / GRA)", "A", "C", "I", "C", "I", "A"),
        ("PPC Loom Allocation & ETA", "I", "R / A", "R", "I", "C", "I"),
        ("Yarn Procurement & PO", "C", "C", "I", "I", "R / A", "C"),
        ("Gate Inward & Store GRN", "I", "I", "I", "C", "R / A", "I"),
        ("Job Work Issuance (DC)", "C", "C", "R / A", "I", "R", "I"),
        ("Job Work Receipt & Wastage", "I", "I", "R", "R", "R / A", "C"),
        ("4-Point Grey Inspection", "I", "I", "C", "R / A", "I", "I"),
        ("Packing Slip & Barcoding", "I", "I", "I", "C", "R / A", "I"),
        ("Sales Invoicing & E-Way", "C", "I", "I", "I", "C", "R / A"),
        ("CubeBook Financial Sync", "I", "I", "I", "I", "I", "R / A"),
        ("HR & Payroll Run", "I", "I", "I", "I", "I", "R / A")
    ]
    for idx, row_data in enumerate(raci_data):
        row = raci_table.add_row()
        format_table_row(row, r_widths, row_data, is_even=(idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 8. DATABASE SCHEMA & KEY ENTITIES OVERVIEW
    # -------------------------------------------------------------
    add_heading_1("7. Database Schema & Key Entities Reference")
    add_paragraph(
        "Handloom ERP features over 60 relational models designed to maintain referential integrity and support rapid analytics. "
        "The following table summarizes the primary database entities:"
    )
    
    schema_table = doc.add_table(rows=1, cols=4)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(schema_table)
    sc_widths = [Inches(1.5), Inches(1.8), Inches(1.5), Inches(1.7)]
    format_table_header(schema_table.rows[0], sc_widths, ["Module Domain", "Table Name", "Primary / Foreign Keys", "Business Function"])
    
    schema_data = [
        ("Parties & Masters", "parties, party_addresses, sub_masters", "id, party_id, entity, code", "Customer, vendor, and sub-master definitions."),
        ("Design & Pre-Prod", "textile_designs, costing_sheets, costing_yarn_lines", "id, design_id, costing_no", "CAD fabric specs, AI color palettes, itemized costs."),
        ("Orders & Merchandising", "buyer_orders, buyer_order_items, proforma_invoices", "id, buyer_order_no, pi_number", "Customer sales contracts, delivery milestones, PIs."),
        ("PPC & Scheduling", "ppc_looms, ppc_schedules, ppc_problems", "id, loom_id, order_id", "Machine allocation, efficiency tracking, downtime logs."),
        ("Yarn Procurement", "yarn_purchase_orders, yarn_inwards", "id, po_number, inward_no", "Raw yarn procurement contracts and gate GRN weighments."),
        ("Job Work Pipeline", "yarn_dyeing_pos, warping_sizing_pos, weaving_pos", "id, po_number, party_id", "External processing contracts for dyeing, sizing, and weaving."),
        ("Cloth & Quality", "cloth_inwards, on_table_checkings, finished_fabrics", "id, piece_no, roll_no", "Woven greige receipts, 4-point defect grading records."),
        ("Warehouse & Stores", "godowns, racks, store_items, store_grns, store_issues", "id, godown_id, item_id", "Multi-warehouse stock bins, departmental stores consumption."),
        ("Sales & Dispatch", "packing_slips, goods_releases, sales_invoices, eway_bills", "id, invoice_no, eway_bill_no", "Finished goods packing, tax invoices, and E-Way bills."),
        ("Fleet Logistics", "vehicles, drivers, routes, trips, fuel_entries", "id, vehicle_id, trip_id", "Truck assignments, delivery routes, fuel expense slips."),
        ("HR & Staff", "employees, departments, designations, attendance, payroll", "id, employee_code", "Staff profiles, biometric attendance, monthly pay slips."),
        ("CubeBook Finance", "companies, ledger_groups, ledgers, vouchers, voucher_entries", "id, company_id, voucher_no", "Double-entry accounting, financial statements, GST."),
        ("Security & Auditing", "users, roles, permissions, log_reports, notifications", "id, user_id, timestamp", "RBAC access governance, audit trails, and alerts.")
    ]
    for idx, row_data in enumerate(schema_data):
        row = schema_table.add_row()
        format_table_row(row, sc_widths, row_data, is_even=(idx % 2 == 1))

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # -------------------------------------------------------------
    # 9. SUMMARY & CONCLUSION
    # -------------------------------------------------------------
    add_heading_1("8. Summary & Conclusion")
    add_paragraph(
        "Handloom ERP represents a modern, comprehensive digital solution engineered specifically to resolve the unique operational, "
        "technical, and financial challenges of textile and handloom manufacturing enterprises. By seamlessly unifying AI-powered design engineering, "
        "merchandising orders, real-time loom planning, multi-stage job work tracking, on-table 4-point quality inspection, warehouse and fleet logistics, "
        "and double-entry accounting in CubeBook, the platform eliminates operational blind spots and protects enterprise margins."
    )
    add_paragraph(
        "With granular Role-Based Access Control (RBAC) and automated audit logging on every transaction, Handloom ERP guarantees complete compliance, "
        "operational transparency, and audit readiness across every department."
    )

    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")
    print(f"File size: {os.path.getsize(output_path)} bytes")

if __name__ == "__main__":
    out_dir = r"c:\Users\user\Documents\Cube Ai Solutions Project\Handloom_ERP"
    out_file = os.path.join(out_dir, "Handloom_ERP_Complete_Project_Workflow.docx")
    build_workflow_document(out_file)
