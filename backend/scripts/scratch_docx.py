# Let's write a python script to search the docx file paragraphs for the word "Weft" and print them.
import docx

doc = docx.Document("/home/cubeai/Videos/Dinesh Export Textile /Dinesh_Textile(may 26)/WEAVING_YARN_CALCULATION_GUIDE.docx")

for idx, para in enumerate(doc.paragraphs):
    if "weft" in para.text.lower():
        print(f"P{idx}: {para.text}")

for table_idx, table in enumerate(doc.tables):
    for row_idx, row in enumerate(table.rows):
        row_text = [cell.text.strip() for cell in row.cells]
        if any("weft" in text.lower() for text in row_text):
            print(f"T{table_idx}R{row_idx}: {row_text}")
