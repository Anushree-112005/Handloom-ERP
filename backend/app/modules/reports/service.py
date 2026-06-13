"""
Report Generation Service — queries real ERP data and generates PDF/Excel/CSV.

Uses deterministic SQLAlchemy queries based on the REPORT_REGISTRY whitelist.
Never executes arbitrary SQL.
"""
import os
import io
import csv
import logging
from datetime import datetime, timezone
from typing import Optional
from pathlib import Path

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.modules.reports.registry import REPORT_REGISTRY, get_report_definition
from app.models.report_job import ReportJob

# Import all queryable models
from app.models.buyer_order import BuyerOrder
from app.models.sales_invoice import SalesInvoice
from app.models.goods_release import GoodsRelease
from app.models.packing_slip import PackingSlip
from app.models.yarn_purchase import YarnPurchaseOrder
from app.models.cloth import ClothInward, ClothDelivery
from app.models.finished_fabric import FinishedFabricInward
from app.models.employee import Employee
from app.models.log_report import LogReport

logger = logging.getLogger(__name__)

# Map model name strings to actual SQLAlchemy model classes
MODEL_MAP = {
    "BuyerOrder": BuyerOrder,
    "SalesInvoice": SalesInvoice,
    "GoodsRelease": GoodsRelease,
    "PackingSlip": PackingSlip,
    "YarnPurchaseOrder": YarnPurchaseOrder,
    "ClothInward": ClothInward,
    "ClothDelivery": ClothDelivery,
    "FinishedFabricInward": FinishedFabricInward,
    "Employee": Employee,
    "LogReport": LogReport,
}

# Date column mapping per model for date-range filtering
DATE_COLUMN_MAP = {
    "BuyerOrder": "order_date",
    "SalesInvoice": "invoice_date",
    "GoodsRelease": "gra_date",
    "PackingSlip": "slip_date",
    "YarnPurchaseOrder": "po_date",
    "ClothInward": "inw_date",
    "ClothDelivery": "dc_date",
    "FinishedFabricInward": "inv_date",
    "Employee": "created_at",
    "LogReport": "created_at",
}

REPORTS_DIR = str(Path(__file__).parent.parent.parent.parent / "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)


async def query_report_data(db: AsyncSession, report_id: str, filters: dict) -> list[dict]:
    """Execute a deterministic query for the given report_id with applied filters."""
    defn = get_report_definition(report_id)
    if not defn:
        raise ValueError(f"Unknown report: {report_id}")

    model_name = defn["model"]
    model_cls = MODEL_MAP.get(model_name)
    if not model_cls:
        raise ValueError(f"Model not configured: {model_name}")

    stmt = select(model_cls)

    # Apply date-range filters
    date_col_name = DATE_COLUMN_MAP.get(model_name)
    if date_col_name and hasattr(model_cls, date_col_name):
        date_col = getattr(model_cls, date_col_name)
        if filters.get("from_date"):
            stmt = stmt.where(date_col >= filters["from_date"])
        if filters.get("to_date"):
            stmt = stmt.where(date_col <= filters["to_date"])

    # Apply party filter
    if filters.get("party") and hasattr(model_cls, "party_name"):
        stmt = stmt.where(model_cls.party_name.ilike(f"%{filters['party']}%"))

    # Apply status filter
    if filters.get("status") and hasattr(model_cls, "status"):
        stmt = stmt.where(model_cls.status == filters["status"])

    # Apply user filter (for log reports)
    if filters.get("user") and hasattr(model_cls, "user_name"):
        stmt = stmt.where(model_cls.user_name.ilike(f"%{filters['user']}%"))

    # Apply department filter (for employees)
    if filters.get("department") and hasattr(model_cls, "department"):
        stmt = stmt.where(model_cls.department == filters["department"])

    # Order by primary date column descending, fall back to id
    if date_col_name and hasattr(model_cls, date_col_name):
        stmt = stmt.order_by(getattr(model_cls, date_col_name).desc())
    else:
        stmt = stmt.order_by(model_cls.id.desc())

    # Limit to 5000 rows max for safety
    stmt = stmt.limit(5000)

    result = await db.execute(stmt)
    rows = result.scalars().all()

    # Convert to dicts using the report column keys
    columns = defn["columns"]
    data = []
    for row in rows:
        row_dict = {}
        for col in columns:
            val = getattr(row, col["key"], None)
            if val is None:
                val = "-"
            elif isinstance(val, datetime):
                val = val.strftime("%Y-%m-%d %H:%M")
            elif hasattr(val, "isoformat"):  # date objects
                val = val.isoformat()
            else:
                val = str(val)
            row_dict[col["key"]] = val
        data.append(row_dict)

    return data


def generate_csv_bytes(report_id: str, data: list[dict]) -> bytes:
    """Generate CSV content from report data."""
    defn = get_report_definition(report_id)
    if not defn:
        raise ValueError(f"Unknown report: {report_id}")
    columns = defn["columns"]

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([c["label"] for c in columns])
    for row in data:
        writer.writerow([row.get(c["key"], "") for c in columns])

    return output.getvalue().encode("utf-8")


def generate_excel_bytes(report_id: str, data: list[dict]) -> bytes:
    """Generate Excel (XLSX) content from report data."""
    try:
        import openpyxl
    except ImportError:
        # Fallback: return CSV if openpyxl not installed
        logger.warning("openpyxl not installed, falling back to CSV for excel format")
        return generate_csv_bytes(report_id, data)

    defn = get_report_definition(report_id)
    if not defn:
        raise ValueError(f"Unknown report: {report_id}")
    columns = defn["columns"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = defn["title"][:31]  # Excel sheet name max 31 chars

    # Header row with styling
    from openpyxl.styles import Font, PatternFill, Alignment
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

    for col_idx, col_def in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_def["label"])
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, col_def in enumerate(columns, 1):
            ws.cell(row=row_idx, column=col_idx, value=row_data.get(col_def["key"], ""))

    # Auto-width columns
    for col_idx, col_def in enumerate(columns, 1):
        max_len = len(col_def["label"])
        for row_data in data[:100]:
            val_len = len(str(row_data.get(col_def["key"], "")))
            if val_len > max_len:
                max_len = val_len
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = min(max_len + 4, 40)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_pdf_bytes(report_id: str, data: list[dict]) -> bytes:
    """Generate PDF content using WeasyPrint with Jinja2 template."""
    import weasyprint
    from jinja2 import Environment, BaseLoader

    defn = get_report_definition(report_id)
    if not defn:
        raise ValueError(f"Unknown report: {report_id}")
    columns = defn["columns"]

    html_template = """
    <!DOCTYPE html>
    <html>
    <head>
      <meta charset="utf-8">
      <style>
        @page { size: A4 landscape; margin: 15mm; }
        body { font-family: 'Helvetica Neue', Arial, sans-serif; font-size: 10px; color: #1e293b; }
        .header { background: linear-gradient(135deg, #4f46e5, #4338ca); color: white; padding: 16px 24px; border-radius: 8px; margin-bottom: 16px; }
        .header h1 { font-size: 18px; margin: 0 0 4px 0; }
        .header p { font-size: 11px; margin: 0; opacity: 0.85; }
        .meta { display: flex; justify-content: space-between; margin-bottom: 12px; font-size: 10px; color: #64748b; }
        table { width: 100%; border-collapse: collapse; }
        th { background: #f1f5f9; padding: 8px 10px; text-align: left; font-size: 9px; text-transform: uppercase; letter-spacing: 0.5px; color: #475569; border-bottom: 2px solid #e2e8f0; }
        td { padding: 7px 10px; border-bottom: 1px solid #e2e8f0; font-size: 10px; }
        tr:nth-child(even) { background: #f8fafc; }
        .footer { margin-top: 16px; text-align: center; font-size: 9px; color: #94a3b8; }
        .total-row { font-weight: bold; background: #eef2ff !important; }
      </style>
    </head>
    <body>
      <div class="header">
        <h1>DINESH EXPORTS — {{ title }}</h1>
        <p>Generated on {{ timestamp }}</p>
      </div>
      <div class="meta">
        <span>Total Records: {{ total_rows }}</span>
        <span>{{ filter_summary }}</span>
      </div>
      <table>
        <thead>
          <tr>
            <th>#</th>
            {% for col in columns %}<th>{{ col.label }}</th>{% endfor %}
          </tr>
        </thead>
        <tbody>
          {% for row in data %}
          <tr>
            <td>{{ loop.index }}</td>
            {% for col in columns %}<td>{{ row[col.key] }}</td>{% endfor %}
          </tr>
          {% endfor %}
          {% if data|length == 0 %}
          <tr><td colspan="{{ columns|length + 1 }}" style="text-align:center; padding:24px; color:#94a3b8;">No data found for the applied filters.</td></tr>
          {% endif %}
        </tbody>
      </table>
      <div class="footer">
        Dinesh Exports ERP &bull; Confidential &bull; {{ timestamp }}
      </div>
    </body>
    </html>
    """

    env = Environment(loader=BaseLoader())
    template = env.from_string(html_template)

    filter_parts = []
    # We don't have filters passed here directly, but the data is already filtered
    filter_summary = " | ".join(filter_parts) if filter_parts else "All records"

    html_out = template.render(
        title=defn["title"],
        timestamp=datetime.now().strftime("%d %b %Y, %H:%M"),
        total_rows=len(data),
        filter_summary=filter_summary,
        columns=columns,
        data=data,
    )

    pdf_bytes = weasyprint.HTML(string=html_out).write_pdf()
    return pdf_bytes


async def cleanup_old_reports(db: AsyncSession):
    """Clean up report files and database jobs older than 24 hours to prevent disk leak."""
    try:
        from datetime import datetime, timedelta, timezone
        cutoff = datetime.now(timezone.utc) - timedelta(hours=24)
        stmt = select(ReportJob).where(ReportJob.created_at < cutoff)
        result = await db.execute(stmt)
        old_jobs = result.scalars().all()
        for job in old_jobs:
            if job.file_path and os.path.exists(job.file_path):
                try:
                    os.remove(job.file_path)
                except Exception as e:
                    logger.warning(f"Failed to delete old report file {job.file_path}: {e}")
            await db.delete(job)
        await db.commit()
    except Exception as e:
        logger.error(f"Error during report cleanup: {e}")


async def generate_report(
    db: AsyncSession,
    report_id: str,
    fmt: str,
    filters: dict,
    user_id: str,
) -> ReportJob:
    """
    Full report generation pipeline:
    1. Clean up expired report files (older than 24 hours)
    2. Validate report_id against registry
    3. Query ERP data
    4. Generate file in requested format (offloaded to threadpool)
    5. Save to disk (offloaded to threadpool)
    6. Create and return ReportJob record
    """
    # Run cleanup first
    await cleanup_old_reports(db)

    defn = get_report_definition(report_id)
    if not defn:
        raise ValueError(f"Unknown report: {report_id}")

    if fmt not in defn["formats"]:
        raise ValueError(f"Format '{fmt}' not supported for {report_id}. Allowed: {defn['formats']}")

    # Create job record
    job = ReportJob(
        report_id=report_id,
        title=defn["title"],
        format=fmt,
        filters=filters,
        status="processing",
        user_id=user_id,
    )
    db.add(job)
    await db.flush()  # get job.id

    try:
        # Query data
        data = await query_report_data(db, report_id, filters)

        # Generate file bytes (offload CPU-bound format generation)
        import asyncio
        if fmt == "csv":
            file_bytes = await asyncio.to_thread(generate_csv_bytes, report_id, data)
            ext = "csv"
            content_type = "text/csv"
        elif fmt == "excel":
            file_bytes = await asyncio.to_thread(generate_excel_bytes, report_id, data)
            ext = "xlsx"
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        elif fmt == "pdf":
            file_bytes = await asyncio.to_thread(generate_pdf_bytes, report_id, data)
            ext = "pdf"
            content_type = "application/pdf"
        else:
            raise ValueError(f"Unsupported format: {fmt}")

        # Save to disk (offload blocking I/O)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        safe_title = defn["title"].replace(" ", "_").replace("/", "-")
        filename = f"{safe_title}_{timestamp}.{ext}"
        filepath = os.path.join(REPORTS_DIR, filename)

        def save_file_sync():
            with open(filepath, "wb") as f:
                f.write(file_bytes)

        await asyncio.to_thread(save_file_sync)

        # Update job
        job.status = "completed"
        job.file_path = filepath
        job.filename = filename
        job.completed_at = datetime.now(timezone.utc)
        await db.commit()

        logger.info(f"Report generated: {filename} ({len(data)} rows, {len(file_bytes)} bytes)")
        return job

    except Exception as e:
        job.status = "failed"
        job.error = str(e)
        await db.commit()
        logger.error(f"Report generation failed for {report_id}: {e}")
        raise

